#!/usr/bin/env python3
"""Produce the Arabic (RTL) edition of the SIRAJ PowerPoint template.

Language policy (brand.config.json → language_policy): SIRAJ communicates in Arabic.
This takes the template built by tools/build-pptx.py and, for the eleven core slides:
  · translates every text, table cell, chart label and speaker note (13-Presentations/Source/pptx-ar.json);
  · mirrors the layout right-to-left (shape positions, paragraph direction and alignment, table direction);
  · sets Tajawal for Arabic and removes letter-spacing (tracking breaks Arabic joining).
The two Arabic sample slides at the end of the bilingual template are dropped (they are superseded).

Usage: python3 tools/pptx-arabize.py <in.pptx> <out.pptx>
"""
import json, sys, os
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AR = json.load(open(os.path.join(ROOT, '13-Presentations/Source/pptx-ar.json'), encoding='utf-8'))
src, dst = sys.argv[1], sys.argv[2]
prs = Presentation(src)
W = prs.slide_width
CORE = 11
missing = set()


def tr(t):
    k = t.strip()
    if not k:
        return t
    if k in AR:
        return AR[k]
    if any(c.isalpha() and c.isascii() for c in k) and '@' not in k and '.net' not in k:
        missing.add(k)
    return t  # numbers, dates and codes stay as they are


def arabic_run(r):
    rPr = r._r.get_or_add_rPr()
    if rPr.get('spc') is not None:
        del rPr.attrib['spc']
    if rPr.get('cap') is not None:
        del rPr.attrib['cap']
    for tag in ('a:latin', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', 'Tajawal')
    rPr.set('lang', 'ar-OM')


def arabize_frame(tf):
    for p in tf.paragraphs:
        text = ''.join(r.text for r in p.runs)
        if text.strip():
            new = tr(text)
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ''
        for r in p.runs:
            arabic_run(r)
        pPr = p._p.get_or_add_pPr()
        pPr.set('rtl', '1')
        al = p.alignment
        if al in (None, PP_ALIGN.LEFT):
            p.alignment = PP_ALIGN.RIGHT
        elif al == PP_ALIGN.RIGHT:
            p.alignment = PP_ALIGN.LEFT


def mirror(sh):
    if sh.left is None or sh.width is None:
        return
    if sh.shape_type is not None and 'PICTURE' in str(sh.shape_type) and sh.width >= W * 0.98:
        xfrm = sh._element.find('.//' + qn('a:xfrm'))
        xfrm.set('flipH', '1')
        return
    sh.left = int(W - sh.left - sh.width)
    # flip directional shapes (connectors, arrows) so they read right-to-left
    xfrm = sh._element.find('.//' + qn('a:xfrm'))
    if xfrm is not None and sh.shape_type is not None and 'LINE' in str(sh.shape_type):
        xfrm.set('flipH', '0' if xfrm.get('flipH') == '1' else '1')


def walk(shapes, nowrap_short=False):
    for sh in shapes:
        if sh.shape_type is not None and 'GROUP' in str(sh.shape_type):
            mirror(sh)
            continue  # group children move with the group
        mirror(sh)
        if sh.has_text_frame:
            arabize_frame(sh.text_frame)
            if nowrap_short and len(sh.text_frame.text.strip()) <= 18 and '\n' not in sh.text_frame.text.strip():
                sh.text_frame.word_wrap = False
        if getattr(sh, 'has_table', False) and sh.has_table:
            # mirror the table by reversing its columns (renders identically in PowerPoint, Keynote and LibreOffice)
            grid = sh._element.find('.//' + qn('a:tblGrid'))
            cols = list(grid)
            for c in cols: grid.remove(c)
            for c in reversed(cols): grid.append(c)
            for tr_ in sh._element.iter(qn('a:tr')):
                tcs = tr_.findall(qn('a:tc'))
                for tc in tcs: tr_.remove(tc)
                for tc in reversed(tcs): tr_.append(tc)
            for row in sh.table.rows:
                for cell in row.cells:
                    arabize_frame(cell.text_frame)
        if getattr(sh, 'has_chart', False) and sh.has_chart:
            ch = sh.chart  # relabel cached categories and series names in the chart XML
            for el in list(ch._chartSpace.iter(qn('c:v'))) + list(ch._chartSpace.iter(qn('a:t'))):
                if el.text and el.text.strip() in AR:
                    el.text = AR[el.text.strip()]


slides = list(prs.slides)
for i, s in enumerate(slides[:CORE]):
    walk(s.shapes, nowrap_short=(i == 6))  # timeline: short labels never wrap
    if s.has_notes_slide:
        n = s.notes_slide.notes_text_frame
        n.text = tr(n.text)

# drop the superseded Arabic sample slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst)[CORE:]:
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)

prs.save(dst)
if missing:
    print('untranslated:', *sorted(missing), sep='\n  ')
print('arabic pptx', dst, f'({CORE} slides)')
