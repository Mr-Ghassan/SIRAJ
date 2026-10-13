#!/usr/bin/env python3
"""SIRAJ Presentation v1.0 — the company & project profile (Arabic, RTL), editable PowerPoint.

    python3 tools/build-presentation-v1.py

Writes 13-Presentations/PPTX/SIRAJ_Presentation_Siraj_AR_v1.0.pptx, then (LibreOffice) the PDF and PNG previews in
13-Presentations/Exports/. Real project only: approved renders and Silk & Light key visuals; contact = phone + Instagram.
Text boxes are native and editable (Tajawal / Jost must be installed on the presenting machine).
"""
import json
import os
import shutil
import subprocess

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
OUTD = os.path.join(ROOT, '13-Presentations')
PPTX = os.path.join(OUTD, 'PPTX', 'SIRAJ_Presentation_Siraj_AR_v1.0.pptx')
EXP = os.path.join(OUTD, 'Exports')
CFG = json.load(open(os.path.join(ROOT, '00-Brand-Foundation', 'brand.config.json')))
PHONE, HANDLE = CFG['contact']['phone_display'], CFG['contact']['social_handle']   # «7771 4404»
WEB = CFG['contact']['website_display']  # always WWW.SIRAJPROPERTIES.NET (client rule 2026-10-13)

HEX = dict(maroon='602733', deep='2C0F12', shadow='170608', stone='AA9180', sand='D6C0B2', ivory='F7F2ED',
           linen='EFE7E0', graphite='3B3436', ash='6B625E', ondark='E9DDD4')
C = {k: RGBColor.from_string(v) for k, v in HEX.items()}
PX = 6350  # 1 px of the 1920×1080 design grid in EMU
M = 120    # side margin (px)

KV = lambda n, r='4x5': os.path.join(ROOT, '14-Property-Marketing/12-Project-Siraj/Source/KeyVisuals', f'SIRAJ_Project-Siraj_KV_{n}_{r}_v1.0.jpg')
CL = lambda n: os.path.join(ROOT, '14-Property-Marketing/12-Project-Siraj/Source/Clean', f'SIRAJ_Project-Siraj_{n}_Clean_v1.0.jpg')
SILK = lambda f: os.path.join(ROOT, '06-SIRAJ-Silk/Stills', f, f'SIRAJ_Silk_{f}_16x9_v1.0.jpg')
LOGO = lambda v: os.path.join(ROOT, '01-Logos/PNG', f'SIRAJ_Logo_{v}_v1.0_2048px.png')

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
BLANK = prs.slide_layouts[6]


def px(v):
    return Emu(int(v * PX))


def bg(slide, color=None, image=None, crop=None):
    if image:
        pic = slide.shapes.add_picture(image, 0, 0, prs.slide_width, prs.slide_height)
        if crop:
            pic.crop_top, pic.crop_bottom = crop
        return pic
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = C[color]


def shade(slide, top=0, h=1080, color='shadow', alpha=55):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, px(top), prs.slide_width, px(h))
    s.fill.solid(); s.fill.fore_color.rgb = C[color]; s.line.fill.background()
    sf = s.fill._xPr.find(qn('a:solidFill'))[0]
    a = etree.SubElement(sf, qn('a:alpha')); a.set('val', str(alpha * 1000))
    return s


def text(slide, x, y, w, h, runs, size=28, color='ivory', align='r', anchor='top', lh=1.5, font='Tajawal'):
    """runs: list of paragraphs; each paragraph = list of (text, weight) or a str."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}[anchor]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'r': PP_ALIGN.RIGHT, 'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER}[align]
        p.line_spacing = lh
        pPr = p._p.get_or_add_pPr(); pPr.set('rtl', '1' if (align != 'l' and font == 'Tajawal') else '0')
        for t, wgt in ([(para, 400)] if isinstance(para, str) else para):
            r = p.add_run(); r.text = t
            f = r.font; f.size = Pt(size * 0.5); f.color.rgb = C[color]
            name = {200: 'Tajawal ExtraLight', 300: 'Tajawal Light', 400: 'Tajawal', 500: 'Tajawal Medium', 700: 'Tajawal', 800: 'Tajawal ExtraBold'}[wgt] if font == 'Tajawal' else font
            f.name = name; f.bold = wgt == 700
            rPr = r._r.get_or_add_rPr()
            for tag in ('a:cs', 'a:latin'):
                el = rPr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rPr, qn(tag))
                el.set('typeface', name)
    return tb


def eyebrow(slide, x, y, w, t, color='sand'):
    text(slide, x, y, w, 40, [[(t, 500)]], size=24, color=color)
    d = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, px(x + w + 14), px(y + 12), px(11), px(11))
    d.fill.solid(); d.fill.fore_color.rgb = C['stone']; d.line.fill.background()


def logo(slide, variant='Sand', w=210, x=None, y=1080 - 96 - 52):
    x = 1920 - M - w if x is None else x  # bottom-right in Arabic layouts
    slide.shapes.add_picture(LOGO(variant), px(x), px(y - w * 0.35 / 2), px(w))


def jost_w(t, size, wgt=500):
    from PIL import ImageFont
    f = ImageFont.truetype(os.path.join(ROOT, '03-Typography/fonts/Jost', 'Jost-Medium.ttf' if wgt >= 500 else 'Jost-Regular.ttf'), size)
    return f.getlength(t)


def ig_png(color):
    """The SIRAJ Instagram glyph (07-Icon-System) as a PNG — it replaces «@» everywhere (client rule 2026-10-13)."""
    from PIL import Image, ImageDraw
    out = os.path.join(EXP, '.lo', f'ig_{color}.png'); os.makedirs(os.path.dirname(out), exist_ok=True)
    k = 16; S = 24 * k; im = Image.new('RGBA', (S, S), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    col = tuple(int(HEX[color][i:i + 2], 16) for i in (0, 2, 4)) + (255,); sw = round(1.8 * k)
    d.rectangle([3.5 * k - sw / 2, 3.5 * k - sw / 2, 20.5 * k + sw / 2, 20.5 * k + sw / 2], outline=col, width=sw)
    d.ellipse([7.6 * k - sw / 2, 7.6 * k - sw / 2, 16.4 * k + sw / 2, 16.4 * k + sw / 2], outline=col, width=sw)
    d.polygon([(17 * k, 5.6 * k), (18.4 * k, 7 * k), (17 * k, 8.4 * k), (15.6 * k, 7 * k)], fill=col)
    im.resize((96, 96), Image.LANCZOS).save(out)
    return out


def ig_line(slide, x, y, size, color, align='l', w=None, wgt=500):
    """[Instagram glyph] HANDLE, left- or centre-aligned at x (centre: x = centre point)."""
    icon = size * 0.92; gap = size * 0.4
    tw = jost_w(HANDLE, size, wgt) + 4
    total = icon + gap + tw
    x0 = x - total / 2 if align == 'c' else x
    top = y + size * 1.12 - icon / 2
    slide.shapes.add_picture(ig_png(color), px(x0), px(top), px(icon), px(icon))
    text(slide, x0 + icon + gap, y, tw + 40, size * 1.6, [[(HANDLE, wgt)]], size=size, color=color, align='l', font='Jost')


def contact(slide, color='sand', y=1080 - 96 - 30):
    # posts-style contact line: phone · [Instagram glyph] HANDLE (never «@»)
    t = f'{PHONE}   ·   '
    text(slide, M, y, 700, 40, [[(t, 500)]], size=22, color=color, align='l', font='Jost')
    ig_line(slide, M + jost_w(t, 22), y, 22, color)


def picture(slide, path, x, y, w, h):
    from PIL import Image
    im = Image.open(path); iw, ih = im.size
    pic = slide.shapes.add_picture(path, px(x), px(y), px(w), px(h))
    r_box, r_img = w / h, iw / ih
    if r_img > r_box:
        c = (1 - r_box / r_img) / 2; pic.crop_left = pic.crop_right = c
    else:
        c = (1 - r_img / r_box) / 2; pic.crop_top = pic.crop_bottom = c
    return pic


# 1 · cover
s = prs.slides.add_slide(BLANK); bg(s, image=SILK('Drape-Lumen')); shade(s, alpha=35)
s.shapes.add_picture(LOGO('Sand'), px(960 - 300), px(400), px(600))
text(s, 360, 640, 1200, 60, [[('سراج للتطوير العقاري', 400)]], size=40, color='ondark', align='c')
# 2 · who we are
s = prs.slides.add_slide(BLANK); bg(s, 'ivory')
picture(s, CL('Corner_Cloudy-Day'), 0, 0, 760, 1080)
eyebrow(s, 860, 150, 940, 'من نحن', 'maroon')
text(s, 860, 230, 940, 360, [[('نطوّر مشاريع عقارية', 200)], [('مدروسة في مسقط.', 800)]], size=84, color='deep', lh=1.36)
text(s, 860, 520, 940, 300, ['نختار المكان بعناية، ونصمّم بالضوء والمادة والتفاصيل، ونبني ما يبقى.'], size=34, color='graphite', lh=1.65)
logo(s, 'Primary'); contact(s, 'ash')
# 3 · the name
s = prs.slides.add_slide(BLANK); bg(s, image=SILK('Drape-Deep')); shade(s, alpha=30)
eyebrow(s, M, 300, 1680, 'معنى الاسم')
text(s, M, 380, 1680, 320, [[('سِراج', 800)], [('المصباح المضيء.', 200)]], size=120, color='ivory', lh=1.32)
logo(s); contact(s)
# 4 · what we believe
s = prs.slides.add_slide(BLANK); bg(s, 'linen')
eyebrow(s, M, 120, 1680, 'ما نؤمن به', 'maroon')
cols = [('الضوء', 'الضوء جزءٌ من العمارة، لا زينةٌ تُضاف إليها.'), ('المادة', 'موادٌ صادقة، تبقى جميلةً مع الزمن.'), ('التفاصيل', 'مدروسة في كل تفاصيلها، من الواجهة إلى المقبض.')]
for i, (h, b) in enumerate(cols):
    x = 1920 - M - (i + 1) * 540 + 30
    text(s, x, 300, 500, 120, [[(h, 800)]], size=64, color='maroon', lh=1.36)
    text(s, x, 430, 500, 300, [b], size=32, color='graphite', lh=1.65)
logo(s, 'Primary'); contact(s, 'ash')
# 5 · the project
s = prs.slides.add_slide(BLANK); bg(s, 'deep')
picture(s, CL('Front_Blue-Hour'), 1920 - 820, 0, 820, 1080)
eyebrow(s, M, 160, 900, 'المشروع')
text(s, M, 250, 900, 190, [[('مسقط', 800)]], size=120, color='ivory', lh=1.36)   # project-name lettering postponed (client 2026-10-13)
text(s, M, 470, 900, 120, [[('شقق سكنية ومحلات تجارية.', 300)]], size=48, color='ondark', lh=1.5)
contact(s); s.shapes.add_picture(LOGO('Sand'), px(M), px(1080 - 96 - 140), px(210))
# 6 · gallery
s = prs.slides.add_slide(BLANK); bg(s, 'ivory')
eyebrow(s, M, 90, 1680, 'تصوّرات المشروع', 'maroon')
names = ['Corner_Cloudy-Day', 'Front_Cloudy-Day', 'Side_Cloudy-Day', 'Corner_Blue-Hour', 'Front_Blue-Hour', 'Side_Blue-Hour']
for i, n in enumerate(names):
    picture(s, CL(n), 1920 - M - (i + 1) * 280 + 12, 170, 268, 700)
logo(s, 'Primary'); contact(s, 'ash')
# 7–9 · key visuals
for kv, a, b in [('Lantern', 'في قلب المدينة،', 'بيتٌ من ضوء.'), ('Ribbon', 'حين يحلّ المساء،', 'يُضيء سراج.'), ('Detail', 'مدروسة', 'في كل تفاصيلها.')]:
    s = prs.slides.add_slide(BLANK); bg(s, 'shadow')
    picture(s, KV(kv), 1920 - 864, 0, 864, 1080)
    order = [(a, 200), (b, 800)] if kv != 'Detail' else [(a, 800), (b, 200)]
    text(s, M, 330, 860, 400, [[order[0]], [order[1]]], size=92, color='ivory', lh=1.36)
    contact(s); s.shapes.add_picture(LOGO('Sand'), px(M), px(1080 - 96 - 140), px(210))
# 10 · end card
s = prs.slides.add_slide(BLANK); bg(s, image=SILK('Drape-Deep')); shade(s, alpha=45)
s.shapes.add_picture(LOGO('Sand'), px(960 - 240), px(360), px(480))
text(s, 360, 610, 1200, 60, [[(PHONE, 500)]], size=44, color='ivory', align='c', font='Jost')
ig_line(s, 960, 680, 34, 'sand', align='c', wgt=400)
text(s, 360, 744, 1200, 60, [[(WEB, 400)]], size=28, color='sand', align='c', font='Jost')

os.makedirs(os.path.dirname(PPTX), exist_ok=True)
prs.save(PPTX)
print('pptx', os.path.relpath(PPTX, ROOT))

# PDF + PNG previews via LibreOffice
os.makedirs(EXP, exist_ok=True)
tmp = os.path.join(EXP, '.lo'); os.makedirs(tmp, exist_ok=True)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, PPTX], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
pdf = os.path.join(EXP, 'SIRAJ_Presentation_Siraj_AR_v1.0.pdf')
shutil.move(os.path.join(tmp, os.path.basename(PPTX).replace('.pptx', '.pdf')), pdf); shutil.rmtree(tmp)
png = os.path.join(EXP, 'PNG'); os.makedirs(png, exist_ok=True)
for f in os.listdir(png):
    if f.startswith('SIRAJ_Presentation_Siraj_AR'):
        os.remove(os.path.join(png, f))
subprocess.run(['pdftoppm', '-r', '72', '-png', pdf, os.path.join(png, 'SIRAJ_Presentation_Siraj_AR_v1.0_slide')], check=True)
print('pdf', os.path.relpath(pdf, ROOT))
