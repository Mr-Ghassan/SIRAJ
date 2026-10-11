#!/usr/bin/env python3
"""SIRAJ Presentation Template v1.0 — editable PowerPoint built with python-pptx.

    python3 tools/build-pptx.py

Writes 13-Presentations/PPTX/SIRAJ_Presentation-Template_v1.0.pptx

Everything is native and editable: text boxes (Jost / Tajawal font names — the fonts must be
installed on the presenting machine), shapes, a native table and native charts in brand colours.
Grid: 1920 x 1080 px design grid mapped onto a 13.333 x 7.5 in slide (1 px = 6350 EMU = 0.5 pt).
"""
import copy
import os
import tempfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
OUT = os.path.join(ROOT, '13-Presentations', 'PPTX', 'SIRAJ_Presentation-Template_v1.0.pptx')
TMP = tempfile.mkdtemp(prefix='siraj-pptx-')

# ---------------------------------------------------------------- brand
HEX = dict(maroon='602733', deep='2C0F12', shadow='170608', stone='AA9180', sand='D6C0B2', ivory='F7F2ED',
           linen='EFE7E0', graphite='3B3436', ash='6B625E', ink='140E10', warm='8E837D', hi='C79A9C',
           mid='81404C', ondark='E9DDD4')
C = {k: RGBColor.from_string(v) for k, v in HEX.items()}
F = dict(light='Jost Light', xlight='Jost ExtraLight', reg='Jost', med='Jost Medium',
         ar_light='Tajawal Light', ar='Tajawal', ar_med='Tajawal Medium')

# grid (px)
M, G, TOP, BOT = 120, 32, 96, 96
COL = (1920 - 2 * M - 11 * G) / 12


def px(v):
    return Emu(int(round(v * 6350)))


def colx(i):
    """x of grid column i (1-based)."""
    return M + (i - 1) * (COL + G)


def span(n):
    return n * COL + (n - 1) * G


def asset(*p):
    return os.path.join(ROOT, *p)


# ---------------------------------------------------------------- primitives
def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = C[color]


def text(slide, x, y, w, h, runs, size=26, font='reg', color='graphite', align='l', track=0, caps=False,
         line=None, rtl=False, anchor='t', space_after=0, name=None):
    """Add a text box. `runs` may be a str, or a list of paragraphs; each paragraph a str or list of
    (text, {overrides}) tuples. size in px (converted to pt)."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}[align]
        if line:
            p.line_spacing = line
        if space_after:
            p.space_after = Pt(space_after * 0.5)
        if rtl:
            p._p.get_or_add_pPr().set('rtl', '1')
        items = para if isinstance(para, list) else [(para, {})]
        for t, o in items:
            r = p.add_run()
            r.text = t.upper() if (o.get('caps', caps)) else t
            fnt = r.font
            fnt.size = Pt(o.get('size', size) * 0.5)
            fnt.name = F[o.get('font', font)]
            fnt.color.rgb = C[o.get('color', color)]
            rPr = r._r.get_or_add_rPr()
            tr = o.get('track', track)
            if tr:
                rPr.set('spc', str(int(round(tr * o.get('size', size) * 0.5 * 100))))
            if rtl or o.get('rtl'):
                rPr.set('lang', 'ar-OM')
                cs = rPr.find(qn('a:cs'))
                if cs is None:
                    cs = etree.SubElement(rPr, qn('a:cs'))
                cs.set('typeface', F[o.get('font', font)])
    return tb


def rect(slide, x, y, w, h, color, line=None, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
    if color:
        s.fill.solid()
        s.fill.fore_color.rgb = C[color]
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = C[line[0]]
        s.line.width = Pt(line[1])
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def hair(slide, x, y, w, color='stone', weight=0.75, alpha=None, vertical=False, h=0, dash=None):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x), px(y), px(x + (0 if vertical else w)), px(y + (h if vertical else 0)))
    st = ln._element.find(qn('p:style'))
    if st is not None:                      # drop theme style (its effectRef adds a shadow)
        ln._element.remove(st)
    ln.line.color.rgb = C[color]
    ln.line.width = Pt(weight)
    if dash:
        ln.line.dash_style = dash
    if alpha is not None:
        sf = ln.line._get_or_add_ln().find(qn('a:solidFill'))
        clr = sf[0]
        a = etree.SubElement(clr, qn('a:alpha'))
        a.set('val', str(int(alpha * 100000)))
    return ln


def lumen(slide, cx, cy, size, color='stone', outline=False):
    s = rect(slide, cx - size / 2, cy - size / 2, size, size, None if outline else color,
             line=(color, 1.0) if outline else None)
    s.rotation = 45
    s.name = 'Lumen'
    return s


def gradient_rect(slide, x, y, w, h, stops, angle=0, name='Scrim'):
    """Rectangle with a linear gradient; stops = [(pos 0-1, colorkey, alpha 0-1)]."""
    s = rect(slide, x, y, w, h, 'shadow', name=name)
    spPr = s._element.spPr
    for tag in ('a:solidFill', 'a:noFill', 'a:gradFill'):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    grad = etree.Element(qn('a:gradFill'), rotWithShape='1')
    gs = etree.SubElement(grad, qn('a:gsLst'))
    for pos, col, alpha in stops:
        g = etree.SubElement(gs, qn('a:gs'), pos=str(int(pos * 100000)))
        c = etree.SubElement(g, qn('a:srgbClr'), val=HEX[col])
        etree.SubElement(c, qn('a:alpha'), val=str(int(alpha * 100000)))
    etree.SubElement(grad, qn('a:lin'), ang=str(int(angle * 60000)), scaled='0')
    # insert after prstGeom
    geom = spPr.find(qn('a:prstGeom'))
    geom.addnext(grad)
    return s


def prepared(path, max_w=1920):
    """Downscale large images into a temp copy so the deck stays light."""
    im = Image.open(path)
    if im.width <= max_w:
        return path, im.size
    out = os.path.join(TMP, os.path.basename(path).rsplit('.', 1)[0] + '.jpg')
    if not os.path.exists(out):
        im = im.convert('RGB')
        im.thumbnail((max_w, max_w * 2), Image.LANCZOS)
        im.save(out, quality=88)
    return out, Image.open(out).size


def picture(slide, path, x, y, w, h, fx=0.5, fy=0.5, shape=None, adj=None, name=None, flip=False):
    """Place an image filling the box (object-fit: cover), optionally inside a preset geometry."""
    src, (iw, ih) = prepared(path)
    pic = slide.shapes.add_picture(src, px(x), px(y), px(w), px(h))
    ar_box, ar_img = w / h, iw / ih
    if ar_img > ar_box:
        frac = ar_box / ar_img
        pic.crop_left, pic.crop_right = (1 - frac) * fx, (1 - frac) * (1 - fx)
    else:
        frac = ar_img / ar_box
        pic.crop_top, pic.crop_bottom = (1 - frac) * fy, (1 - frac) * (1 - fy)
    if shape is not None:
        pic.auto_shape_type = shape
        if adj is not None:
            geom = pic._element.spPr.find(qn('a:prstGeom'))
            av = geom.find(qn('a:avLst'))
            if av is None:
                av = etree.SubElement(geom, qn('a:avLst'))
            for i, v in enumerate(adj):
                etree.SubElement(av, qn('a:gd'), name='adj' if len(adj) == 1 else f'adj{i + 1}', fmla=f'val {int(v)}')
    if flip:
        xfrm = pic._element.spPr.find(qn('a:xfrm'))
        xfrm.set('flipH', '1')
    if name:
        pic.name = name
    return pic


def logo(slide, x, y, w, variant='Primary'):
    p = slide.shapes.add_picture(asset('01-Logos', 'PNG', f'SIRAJ_Logo_{variant}_v1.0_1024px.png'), px(x), px(y), width=px(w))
    p.name = 'SIRAJ wordmark'
    return p


def footer(slide, n, section, dark=False, start_x=M, rtl=False):
    col = 'sand' if dark else 'ash'
    y = 1008
    if rtl:
        lumen(slide, 1800 - 4, y + 24, 7)
        text(slide, 1100, y + 14, 684, 24, section, size=17, font='ar_med', color=col, align='r', rtl=True, name='Footer · section')
        logo(slide, M, y + 10, 96, 'Sand' if dark else 'Primary')
        text(slide, M + 120, y + 14, 60, 24, f'{n:02d}', size=14, font='med', color=col, track=0.16, name='Footer · page')
        return
    lumen(slide, start_x + 4, y + 24, 7)
    text(slide, start_x + 20, y + 15, 700, 24, section, size=14, font='med', color=col, track=0.16, caps=True, name='Footer · section')
    text(slide, 1800 - 96 - 28 - 60, y + 15, 60, 24, f'{n:02d}', size=14, font='med', color=col, track=0.16, align='r', name='Footer · page')
    logo(slide, 1800 - 96, y + 10, 96, 'Sand' if dark else 'Primary')


def eyebrow(slide, x, y, w, t, dark=False, rtl=False):
    if rtl:
        return text(slide, x, y, w, 30, t, size=20, font='ar_med', color='sand' if dark else 'maroon', align='r', rtl=True, name='Eyebrow')
    return text(slide, x, y, w, 24, t, size=16, font='med', color='sand' if dark else 'maroon', track=0.16, caps=True, name='Eyebrow')


def title(slide, x, y, w, h, t, dark=False, size=60, rtl=False):
    if rtl:
        return text(slide, x, y, w, h, t, size=64, font='ar_light', color='ivory' if dark else 'ink', align='r', rtl=True, line=1.15, name='Title')
    return text(slide, x, y, w, h, t, size=size, font='light', color='ivory' if dark else 'ink', line=0.98, name='Title')


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ---------------------------------------------------------------- theme (fonts + colours) & layout names
def brand_theme(prs):
    master = prs.slide_master
    for rel in master.part.rels.values():
        if rel.reltype.endswith('/theme'):
            part = rel.target_part
            root = etree.fromstring(part.blob)
            a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
            ns = {'a': a}
            root.set('name', 'SIRAJ')
            cs = root.find('.//a:clrScheme', ns)
            cs.set('name', 'SIRAJ')
            mapping = {'dk1': 'ink', 'lt1': 'ivory', 'dk2': 'deep', 'lt2': 'linen', 'accent1': 'maroon', 'accent2': 'stone',
                       'accent3': 'graphite', 'accent4': 'hi', 'accent5': 'warm', 'accent6': 'deep', 'hlink': 'maroon', 'folHlink': 'mid'}
            for tag, key in mapping.items():
                el = cs.find(f'a:{tag}', ns)
                for ch in list(el):
                    el.remove(ch)
                etree.SubElement(el, f'{{{a}}}srgbClr', val=HEX[key])
            fs = root.find('.//a:fontScheme', ns)
            fs.set('name', 'SIRAJ')
            for tag, latin in (('a:majorFont', F['light']), ('a:minorFont', F['reg'])):
                el = fs.find(tag, ns)
                el.find('a:latin', ns).set('typeface', latin)
                el.find('a:cs', ns).set('typeface', F['ar_light'] if 'major' in tag else F['ar'])
                for f in el.findall('a:font', ns):
                    if f.get('script') == 'Arab':
                        f.set('typeface', F['ar_light'] if 'major' in tag else F['ar'])
            part._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    # brand the master background
    bg = master.background.fill
    bg.solid()
    bg.fore_color.rgb = C['ivory']


# =============================================================================
def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(1920), px(1080)
    brand_theme(prs)
    blank = prs.slide_layouts[6]
    blank.name = 'SIRAJ · Blank'
    n = 0

    def new(bg='ivory'):
        nonlocal n
        n += 1
        s = prs.slides.add_slide(blank)
        set_bg(s, bg)
        return s

    # 01 COVER -----------------------------------------------------------------
    s = new('shadow')
    picture(s, asset('08-Photography', 'Library', 'SIRAJ_Photo_Campaign-Silk-Wrap_v1.0.jpg'), colx(8), TOP, 1920 - colx(8), 1080 - TOP,
            fy=0.4, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, adj=[34000, 0], name='Portal image')
    logo(s, M, TOP, 168, 'Sand')
    lumen(s, M + 4, 512, 8)
    eyebrow(s, M + 22, 500, 700, 'Investor presentation · October 2026', dark=True)
    text(s, M, 552, span(7), 130, 'Noor Residences', size=112, font='light', color='ivory', name='Title')
    text(s, M, 690, span(6), 70, 'مساكن نور', size=48, font='ar_light', color='sand', rtl=True, align='l', name='Arabic name')
    text(s, M, 790, 640, 100, 'A 24-storey tower of 168 residences in Al Khuwair, Muscat.', size=32, font='light', color='ondark', line=1.25)
    text(s, M, 1023, 900, 24, 'SIRAJ Real Estate Development · Confidential', size=14, font='med', color='sand', track=0.16, caps=True)
    notes(s, 'Layout: Cover. Replace the portal image (right) by right-click → Change Picture; the rounded top keeps the SIRAJ Portal shape.')

    # 02 SECTION DIVIDER --------------------------------------------------------
    s = new('shadow')
    picture(s, asset('06-SIRAJ-Silk', 'Stills', 'Drape-Lumen', 'SIRAJ_Silk_Drape-Lumen_16x9_v1.0.jpg'), 0, 0, 1920, 1080, name='Silk background')
    gradient_rect(s, 0, 0, 1920, 1080, [(0, 'shadow', 0.78), (0.38, 'shadow', 0.45), (0.62, 'shadow', 0)], angle=0)
    text(s, M - 8, 400, 600, 220, '01', size=240, font='xlight', color='sand', line=0.8, name='Section number')
    text(s, M, 640, span(7), 130, 'The project', size=112, font='light', color='ivory', name='Title')
    text(s, M, 790, 640, 100, 'Twenty-four storeys shaped by light, on Muscat\'s coastal spine.', size=32, font='light', color='ondark', line=1.25)
    notes(s, 'Layout: Section divider (dark silk). Use silk only on cover, dividers and closing.')

    # 03 TITLE + TEXT (executive summary) ----------------------------------------
    s = new('ivory')
    eyebrow(s, M, TOP, span(5), 'Executive summary')
    title(s, M, TOP + 40, span(5), 220, 'A quiet tower at the centre of Muscat, half sold in seven months.')
    text(s, M, 760, span(5), 130, [[('84', {'size': 120, 'font': 'xlight', 'color': 'ink'}), (' / 168', {'size': 36, 'font': 'light', 'color': 'ash'})]])
    hair(s, M, 924, span(5), color='stone', weight=1.5, alpha=0.35)
    hair(s, M, 924, span(5) / 2, color='maroon', weight=1.5)
    lumen(s, M + span(5) / 2, 924, 14)
    text(s, M, 948, span(5), 24, 'Residences sold · September 2026', size=16, font='med', color='ash', track=0.16, caps=True)
    pts = [('A prime, connected address', 'On Sultan Qaboos Street in Al Khuwair, five minutes from the Grand Mosque and fifteen from the airport.'),
           ('Strong early absorption', '84 of 168 residences sold since launch in March 2026, about twelve each month.'),
           ('Resilient economics', 'OMR 40.0 M gross development value and a projected 27% gross margin on an OMR 29.2 M budget.'),
           ('On programme for Q4 2027', 'Foundations and podium are complete. The structure has reached level 9 of 24.')]
    y0, step = 104, 212
    for i, (h, b) in enumerate(pts):
        y = y0 + i * step
        hair(s, colx(7), y, span(6))
        lumen(s, colx(7) + 6, y + 62, 10)
        text(s, colx(7) + 56, y + 40, span(6) - 56, 44, h, size=32, font='reg', color='ink')
        text(s, colx(7) + 56, y + 92, span(6) - 56, 90, b, size=24, font='reg', color='ash', line=1.2)
    hair(s, colx(7), y0 + 4 * step, span(6))
    footer(s, n, 'Executive summary')
    notes(s, 'Layout: Title + text. Title 56–64 px Jost Light (28–32 pt), body 24–28 px (12–14 pt).')

    # 04 TWO-COLUMN (project overview) ---------------------------------------------
    s = new('ivory')
    picture(s, asset('08-Photography', 'Library', 'SIRAJ_Photo_Residential-Tower-Golden-Hour_v1.0.jpg'), 0, 0, M + span(6), 1080,
            shape=MSO_SHAPE.ROUND_1_RECTANGLE, adj=[30000], name='Image · Silk Corner')
    eyebrow(s, colx(8), TOP, span(5), 'Project overview')
    title(s, colx(8), TOP + 40, span(5), 80, 'Noor Residences')
    text(s, colx(8), 236, span(5), 130, 'A slender tower of warm stone and bronze. Deep loggias temper the afternoon sun and frame long views to the Gulf of Oman.',
         size=24, color='ash', line=1.25)
    specs = [('Location', 'Al Khuwair, Muscat'), ('Typology', 'Residential tower'), ('Floors', '24 + podium'), ('Residences', '168'),
             ('Plot area', '5,170 m²'), ('Gross floor area', '31,500 m²'), ('Saleable area', '23,930 m²'), ('Handover', 'Q4 2027')]
    cw = (span(5) - G) / 2
    for i, (k, v) in enumerate(specs):
        x = colx(8) + (i % 2) * (cw + G)
        y = 548 + (i // 2) * 109
        hair(s, x, y, cw)
        text(s, x, y + 24, cw, 22, k, size=15, font='med', color='ash', track=0.16, caps=True)
        text(s, x, y + 52, cw, 40, v, size=30, font='light', color='ink')
    footer(s, n, '01 · The project', start_x=colx(8))
    notes(s, 'Layout: Two-column — image with one Silk Corner (top-end) + specification grid.')

    # 05 KEY NUMBERS ---------------------------------------------------------------
    s = new('shadow')
    eyebrow(s, M, TOP, span(8), 'Key numbers', dark=True)
    title(s, M, TOP + 40, span(8), 80, 'Noor Residences at a glance', dark=True)
    kpis = [('168', '', 'Residences across 24 floors'), ('50', '%', 'Sold within seven months of launch'), ('40.0', ' M', 'Gross development value · OMR'),
            ('27', '%', 'Projected gross margin'), ('6.4', '%', 'Projected gross rental yield'), ('2027', '', 'Handover · Q4')]
    kw = span(4)
    for i, (v, u, l) in enumerate(kpis):
        x = M + (i % 3) * (kw + G)
        y = 392 + (i // 3) * 296
        hair(s, x, y, kw, color='sand', alpha=0.28)
        text(s, x, y + 36, kw, 150, [[(v, {}), (u, {'size': 40, 'font': 'light', 'color': 'sand'})]], size=132, font='xlight', color='ivory')
        text(s, x, y + 196, kw, 40, l, size=24, color='ondark')
        if i == 1:
            lumen(s, x + kw - 12, y + 64, 18)
    footer(s, n, '02 · Performance', dark=True)
    notes(s, 'Layout: Key numbers (dark). Numerals Jost ExtraLight; units in Sand. Mark the key figure with one Lumen.')

    # 06 CHARTS (native) ---------------------------------------------------------------
    s = new('ivory')
    eyebrow(s, M, TOP, span(8), 'Sales & pricing')
    title(s, M, TOP + 40, span(10), 80, 'Steady absorption at a premium to the market')
    for x, h, u in ((M, 'Sales velocity', 'Residences sold per month, 2026 · 84 to date'),
                    (colx(7), 'Price per m²', 'Average asking price, OMR per m² of saleable area')):
        text(s, x, 292, span(6), 40, h, size=30, color='ink')
        text(s, x, 338, span(6), 30, u, size=20, color='ash')
    # bar chart
    cd = CategoryChartData()
    cd.categories = ['Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep']
    cd.add_series('Residences sold', (17, 12, 9, 11, 10, 12, 13))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, px(M - 16), px(400), px(span(6) + 16), px(500), cd)
    gf.name = 'Chart · Sales velocity'
    style_chart(gf.chart)
    plot = gf.chart.plots[0]
    plot.gap_width = 90
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = C['stone']
    pt = ser.points[0]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = C['maroon']
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.number_format, dl.number_format_is_linked = '0', False
    dl.font.size, dl.font.name, dl.font.color.rgb = Pt(12), F['light'], C['graphite']
    va = gf.chart.value_axis
    va.maximum_scale, va.minimum_scale, va.major_unit = 20, 0, 5
    # line chart
    cd = CategoryChartData()
    cd.categories = ['2021', '2022', '2023', '2024', '2025', '2026']
    cd.add_series('Al Khuwair new-build', (1180, 1230, 1290, 1370, 1460, 1560))
    cd.add_series('Muscat apartments', (760, 780, 805, 840, 885, 930))
    cd.add_series('Noor Residences', (None, None, None, None, None, 1670))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, px(colx(7) - 16), px(400), px(span(6) + 16), px(500), cd)
    gf.name = 'Chart · Price per m²'
    ch = gf.chart
    style_chart(ch)
    va = ch.value_axis
    va.maximum_scale, va.minimum_scale, va.major_unit = 1800, 600, 300
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = '#,##0', False
    plot_layout(ch, 0.09, 0.06, 0.66, 0.82)
    for i, (key, w, msz, mstyle) in enumerate((('maroon', 2.25, 7, XL_MARKER_STYLE.CIRCLE), ('warm', 1.5, 6, XL_MARKER_STYLE.CIRCLE), ('stone', 0, 13, XL_MARKER_STYLE.DIAMOND))):
        se = ch.plots[0].series[i]
        se.smooth = False
        if w:
            se.format.line.color.rgb = C[key]
            se.format.line.width = Pt(w)
        else:
            se.format.line.fill.background()
        se.marker.style, se.marker.size = mstyle, msz
        se.marker.format.fill.solid()
        se.marker.format.fill.fore_color.rgb = C['stone'] if i == 2 else C['ivory']
        se.marker.format.line.color.rgb = C[key]
        lab = se.points[5].data_label
        lab.has_text_frame = True
        lab.text_frame.text = ['1,560 Al Khuwair', '930 Muscat', '1,670 Noor'][i]
        lab.position = XL_LABEL_POSITION.RIGHT if i < 2 else XL_LABEL_POSITION.ABOVE
        for p in lab.text_frame.paragraphs:
            for r in p.runs:
                r.font.size, r.font.name = Pt(11), F['reg']
                r.font.color.rgb = C['maroon'] if i == 2 else C['ink']
    text(s, M, 952, span(12), 24, 'Source: SIRAJ sales records, Sep 2026. Market series are illustrative demonstration data.', size=15, color='ash')
    footer(s, n, '02 · Performance')
    notes(s, 'Layout: Charts. Native charts — right-click → Edit Data. Brand rules: hairline gridlines, light numerals, direct labels (no legend), Maroon for the highlight, Stone for the rest, Lumen ◆ for the key point.')

    # 07 TIMELINE --------------------------------------------------------------------
    s = new('ivory')
    eyebrow(s, M, TOP, span(7), 'Delivery')
    title(s, M, TOP + 40, span(7), 150, 'On programme for handover in Q4 2027')
    text(s, M, 300, span(7), 40, 'The structure has reached level 9 of 24. Construction is 38% complete.', size=26, color='ash')
    picture(s, asset('08-Photography', 'Library', 'SIRAJ_Photo_Construction-Site-Dawn_v1.0.jpg'), colx(8), TOP, span(5), 300, name='Site photo')
    text(s, colx(8), 412, span(5), 24, 'Site progress · October 2026', size=16, font='med', color='ash', track=0.16, caps=True)
    ly = 760
    hair(s, M, ly, 1680, color='stone', weight=0.75)
    hair(s, M, ly, 910, color='maroon', weight=1.5)
    hair(s, M + 910, ly - 96, 0, color='maroon', weight=0.75, vertical=True, h=96, dash=4)
    text(s, M + 924, ly - 104, 400, 24, 'Today · Oct 2026', size=15, font='med', color='maroon', track=0.16, caps=True)
    items = [('Q2 2025', 'Land acquired', 'Plot secured on Sultan Qaboos Street'), ('Q4 2025', 'Permits & design', 'Building permit and design freeze'),
             ('Q1 2026', 'Sales launch', 'Released to market in March'), ('Q2 2026', 'Foundations', 'Piling and two basement levels'),
             ('Q3 2026', 'Podium', 'Podium slab and amenity shell'), ('Q2 2027', 'Topping out', 'Structure complete to level 24'),
             ('Q3 2027', 'Façade & fit-out', 'Envelope sealed, interiors underway'), ('Q4 2027', 'Handover', 'Keys to the first residents')]
    iw = (1680 - 7 * G) / 8
    for i, (q, h, b) in enumerate(items):
        x = M + i * (iw + G)
        if i < 5:
            lumen(s, x + 8, ly, 16, 'maroon')
        elif i < 7:
            lumen(s, x + 8, ly, 16, 'stone', outline=True)
        else:
            lumen(s, x + 8, ly, 16, 'stone')
        text(s, x, ly + 46, iw, 40, q, size=34, font='light', color='ink')
        text(s, x, ly + 100, iw, 30, h, size=22, font='med', color='ink')
        text(s, x, ly + 136, iw, 60, b, size=19, color='ash', line=1.15)
    footer(s, n, '02 · Performance')
    notes(s, 'Layout: Timeline. Filled Maroon Lumen = complete, outlined = upcoming, Stone = key milestone. Move the dashed Today marker as the programme advances.')

    # 08 IMAGE + CAPTION ------------------------------------------------------------------
    s = new('shadow')
    picture(s, asset('08-Photography', 'Library', 'SIRAJ_Photo_Apartment-Interior-Sea-View_v1.0.jpg'), 0, 0, 1920, 1080, fy=0.6, name='Full-bleed image')
    gradient_rect(s, 0, 0, 1920, 1080, [(0, 'shadow', 0.72), (0.38, 'shadow', 0.4), (0.62, 'shadow', 0)], angle=0)
    gradient_rect(s, 0, 600, 1920, 480, [(0, 'shadow', 0), (1, 'shadow', 0.5)], angle=90)
    eyebrow(s, M, TOP, span(5), 'The residences', dark=True)
    title(s, M, TOP + 40, span(5), 150, 'Every home opens to the light of the Gulf.', dark=True)
    text(s, M, 316, span(6), 24, 'Typical two-bedroom living room · Visualisation', size=16, font='med', color='sand', track=0.16, caps=True)
    footer(s, n, '01 · The project', dark=True)
    notes(s, 'Layout: Full-bleed image + caption. Keep text in the calm, darker third of the photograph.')

    # 09 TABLE ---------------------------------------------------------------------------------
    s = new('ivory')
    eyebrow(s, M, TOP, span(8), 'Financials')
    title(s, M, TOP + 40, span(8), 80, 'Development budget and returns')
    rows = [('Budget line', 'OMR M', 'Share of cost', 'Committed to date'), ('Land', '4.6', '15.8%', '100%'), ('Construction', '19.8', '67.8%', '46%'),
            ('Professional fees', '1.9', '6.5%', '72%'), ('Marketing & sales', '1.2', '4.1%', '38%'), ('Finance & contingency', '1.7', '5.8%', '30%'),
            ('Total development cost', '29.2', '100.0%', '55%')]
    brand_table(s, M, 290, span(8), rows, [0.40, 0.18, 0.21, 0.21])
    text(s, M, 860, span(8), 24, 'Figures in OMR million unless stated. Projections are illustrative and subject to change.', size=15, color='ash')
    ret = [('Gross development value', 'OMR 40.0 M'), ('Gross profit', 'OMR 10.8 M'), ('Gross margin', '27.0%'), ('Project IRR, unlevered', '19.6%'), ('Contracted sales', 'OMR 18.4 M')]
    for i, (k, v) in enumerate(ret):
        y = 290 + i * 128
        hair(s, colx(10), y, span(3))
        text(s, colx(10), y + 22, span(3), 22, k, size=15, font='med', color='ash', track=0.16, caps=True)
        text(s, colx(10), y + 54, span(3), 50, v, size=40, font='light', color='maroon' if i == 2 else 'ink')
    hair(s, colx(10), 290 + 5 * 128, span(3))
    footer(s, n, '03 · Investment')
    notes(s, 'Layout: Data table — native table, no fills, hairline rules, figures right-aligned in light tabular numerals.')

    # 10 QUOTE -------------------------------------------------------------------------------------
    s = new('deep')
    gradient_halo(s, 157, 256, 120)
    lumen(s, 157, 256, 52, 'stone')
    text(s, M, 430, span(9), 260, 'We build slowly, so that the people who live here can live quietly.', size=76, font='light', color='ivory', line=1.05)
    hair(s, M, 712, 48, color='sand', alpha=0.6)
    text(s, M + 68, 700, 1000, 24, 'Chief Executive Officer · SIRAJ Real Estate Development', size=16, font='med', color='sand', track=0.16, caps=True)
    footer(s, n, 'Closing', dark=True)
    notes(s, 'Layout: Quote. One sentence, Jost Light, attribution in caption style.')

    # 11 CLOSING ------------------------------------------------------------------------------------
    s = new('shadow')
    picture(s, asset('06-SIRAJ-Silk', 'Stills', 'Gateway', 'SIRAJ_Silk_Gateway_16x9_v1.0.jpg'), 0, 0, 1920, 1080, name='Silk background')
    logo(s, 960 - 170, 330, 340, 'Sand')
    text(s, 0, 482, 1920, 50, 'Light, built.', size=32, font='light', color='sand', align='c')
    text(s, 0, 614, 1920, 24, 'Book a private viewing at the SIRAJ Sales Centre', size=16, font='med', color='sand', track=0.16, caps=True, align='c')
    text(s, 0, 664, 1920, 36, 'sirajproperties.net   ◆   info@SIRAJPROPERTIES.net   ◆   +968 2400 0000   ◆   Al Khuwair, Muscat', size=22, font='light', color='ondark', align='c')
    text(s, 0, 992, 1920, 24, 'Noor Residences is a demonstration project. All figures are illustrative and do not constitute an offer.', size=15, color='sand', align='c')
    notes(s, 'Layout: Closing / contact. Contact details come from 00-Brand-Foundation/brand.config.json (placeholders until confirmed).')

    # 12 ARABIC COVER (RTL) -----------------------------------------------------------------------------
    s = new('shadow')
    picture(s, asset('08-Photography', 'Library', 'SIRAJ_Photo_Campaign-Silk-Wrap_v1.0.jpg'), 0, TOP, 1920 - colx(8), 1080 - TOP,
            fy=0.4, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, adj=[34000, 0], name='Portal image')
    logo(s, 1800 - 168, TOP, 168, 'Sand')
    eyebrow(s, colx(7), 470, span(6) - 24, 'عرض المستثمرين · أكتوبر 2026', dark=True, rtl=True)
    lumen(s, 1796, 486, 8)
    text(s, colx(7), 520, span(6), 150, 'مساكن نور', size=116, font='ar_light', color='ivory', align='r', rtl=True, name='Title')
    text(s, colx(7), 680, span(6), 60, 'Noor Residences', size=44, font='light', color='sand', align='r')
    text(s, colx(7), 770, span(6), 120, 'برج سكني من 24 طابقاً يضم 168 وحدة سكنية في الخوير، مسقط.', size=32, font='ar_light', color='ondark', align='r', rtl=True, line=1.3)
    text(s, colx(7), 1018, span(6), 30, 'سراج للتطوير العقاري · سرّي', size=17, font='ar_med', color='sand', align='r', rtl=True)
    notes(s, 'Layout: Arabic cover (RTL). Paragraphs are set right-to-left; numbers stay Western (0–9). Arabic is never tracked or uppercased.')

    # 13 ARABIC TITLE + TEXT (RTL) ------------------------------------------------------------------------
    s = new('ivory')
    eyebrow(s, colx(8), TOP, span(5), 'الملخص التنفيذي', rtl=True)
    title(s, colx(8), TOP + 40, span(5), 260, 'برج هادئ في قلب مسقط، بيع نصفه خلال سبعة أشهر.', rtl=True)
    pts = [('عنوان مميّز ومتّصل', 'على شارع السلطان قابوس في الخوير، على بُعد خمس دقائق من الجامع الأكبر وربع ساعة من المطار.'),
           ('إقبال قوي منذ الإطلاق', 'بيعت 84 وحدة من أصل 168 منذ الإطلاق في مارس 2026، بمعدل 12 وحدة شهرياً تقريباً.'),
           ('جدوى اقتصادية متينة', 'قيمة تطوير إجمالية 40.0 مليون ر.ع. وهامش ربح إجمالي متوقع 27% على ميزانية 29.2 مليون ر.ع.'),
           ('التسليم في الربع الرابع 2027', 'اكتملت الأساسات والمنصة، ووصل الهيكل إلى الطابق 9 من 24.')]
    for i, (h, b) in enumerate(pts):
        y = 104 + i * 212
        hair(s, M, y, span(6))
        lumen(s, M + span(6) - 6, y + 64, 10)
        text(s, M, y + 36, span(6) - 56, 50, h, size=33, font='ar_med', color='ink', align='r', rtl=True)
        text(s, M, y + 96, span(6) - 56, 100, b, size=25, font='ar', color='ash', align='r', rtl=True, line=1.35)
    hair(s, M, 104 + 4 * 212, span(6))
    footer(s, n, 'الملخص التنفيذي', rtl=True)
    notes(s, 'Layout: Arabic title + text (RTL). Mirror the grid: title on the right (start), content on the left.')

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.core_properties.title = 'SIRAJ Presentation Template v1.0'
    prs.core_properties.author = 'SIRAJ Real Estate Development'
    prs.core_properties.subject = 'Noor Residences — sample investor presentation (demonstration project)'
    prs.save(OUT)
    print('  pptx', os.path.relpath(OUT, ROOT), f'({n} slides)')


def gradient_halo(slide, cx, cy, r):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(cx - r), px(cy - r), px(2 * r), px(2 * r))
    s.line.fill.background()
    s.shadow.inherit = False
    spPr = s._element.spPr
    for tag in ('a:solidFill', 'a:noFill'):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    grad = etree.Element(qn('a:gradFill'), rotWithShape='1')
    gs = etree.SubElement(grad, qn('a:gsLst'))
    for pos, alpha in ((0, 0.55), (100000, 0)):
        g = etree.SubElement(gs, qn('a:gs'), pos=str(pos))
        c = etree.SubElement(g, qn('a:srgbClr'), val=HEX['sand'])
        etree.SubElement(c, qn('a:alpha'), val=str(int(alpha * 100000)))
    path = etree.SubElement(grad, qn('a:path'), path='circle')
    etree.SubElement(path, qn('a:fillToRect'), l='50000', t='50000', r='50000', b='50000')
    spPr.find(qn('a:prstGeom')).addnext(grad)
    s.name = 'Lumen halo'


def plot_layout(ch, x, y, w, h):
    """Manual inner plot-area layout (fractions of the chart frame) — leaves room for direct labels."""
    pa = ch._chartSpace.find(qn('c:chart')).find(qn('c:plotArea'))
    lay = pa.find(qn('c:layout'))
    if lay is None:
        lay = etree.Element(qn('c:layout'))
        pa.insert(0, lay)
    for el in list(lay):
        lay.remove(el)
    ml = etree.SubElement(lay, qn('c:manualLayout'))
    for tag, val in (('c:layoutTarget', 'inner'), ('c:xMode', 'edge'), ('c:yMode', 'edge'), ('c:x', x), ('c:y', y), ('c:w', w), ('c:h', h)):
        etree.SubElement(ml, qn(tag), val=str(val))


def style_chart(ch):
    ch.has_title = False
    ch.has_legend = False
    ch.font.name, ch.font.size = F['light'], Pt(10)
    ch.font.color.rgb = C['ash']
    va, ca = ch.value_axis, ch.category_axis
    va.has_major_gridlines = True
    gl = va.major_gridlines.format.line
    gl.color.rgb, gl.width = C['stone'], Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.size, va.tick_labels.font.name = Pt(9), F['light']
    va.tick_labels.font.color.rgb = C['ash']
    ca.format.line.color.rgb, ca.format.line.width = C['graphite'], Pt(0.75)
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(9), F['reg']
    ca.tick_labels.font.color.rgb = C['ash']
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    from pptx.enum.chart import XL_TICK_MARK
    ca.major_tick_mark = va.major_tick_mark = XL_TICK_MARK.NONE
    # no chart-area fill or border
    cs = ch._chartSpace
    spPr = cs.find(qn('c:spPr'))
    if spPr is None:
        spPr = etree.SubElement(cs, qn('c:spPr'))
        cs.insert(cs.index(cs.find(qn('c:chart'))) + 1, spPr)
    for el in list(spPr):
        spPr.remove(el)
    etree.SubElement(spPr, qn('a:noFill'))
    ln = etree.SubElement(spPr, qn('a:ln'))
    etree.SubElement(ln, qn('a:noFill'))


def brand_table(slide, x, y, w, rows, colw):
    nr, nc = len(rows), len(rows[0])
    rh = [62] + [78] * (nr - 1)
    gf = slide.shapes.add_table(nr, nc, px(x), px(y), px(w), px(sum(rh)))
    gf.name = 'Table · Budget'
    tbl = gf.table
    # "No Style, No Grid"
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    for j, f in enumerate(colw):
        tbl.columns[j].width = px(w * f)
    for i, row in enumerate(rows):
        tbl.rows[i].height = px(rh[i])
        head, total = i == 0, i == nr - 1
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = 0
            cell.margin_top = cell.margin_bottom = 0
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.text = ''
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            r = p.add_run()
            r.text = val.upper() if head else val
            fnt = r.font
            if head:
                fnt.size, fnt.name, fnt.color.rgb = Pt(7.5), F['med'], C['ash']
                r._r.get_or_add_rPr().set('spc', '120')
            elif j == 0:
                fnt.size, fnt.name, fnt.color.rgb = Pt(12), F['med'] if total else F['reg'], C['ink'] if total else C['graphite']
            else:
                fnt.size, fnt.name, fnt.color.rgb = Pt(13), F['reg'] if total else F['light'], C['ink']
            # borders: bottom hairline (graphite under header and above total)
            tcPr = cell._tc.get_or_add_tcPr()
            for side in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = etree.SubElement(tcPr, qn(side))
                strong = (side == 'a:lnB' and (head or i == nr - 2))
                if side == 'a:lnB' and not total:
                    ln.set('w', str(9525 if strong else 6350))
                    sf = etree.SubElement(ln, qn('a:solidFill'))
                    etree.SubElement(sf, qn('a:srgbClr'), val=HEX['graphite'] if strong else HEX['stone'])
                else:
                    ln.set('w', '0')
                    etree.SubElement(ln, qn('a:noFill'))
            etree.SubElement(tcPr, qn('a:noFill'))


if __name__ == '__main__':
    build()
